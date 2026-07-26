"""Regression tests for P0-6: rca_core.exporter.to_xlsx must preserve
string entries in other_fossils instead of silently dropping them.

REVIEW-2026-07-25 P0-6.
"""
from __future__ import annotations

import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
import io

openpyxl = pytest.importorskip("openpyxl")

from rca_core.exporter import to_xlsx
from openpyxl import load_workbook


def _other_data(*strings):
    return {
        "sections": [], "species_ranges": [], "biozones": [],
        "other_fossils": list(strings), "confidence": 0.5,
    }


def _find_sheet(wb, keyword):
    for s in wb.worksheets:
        if keyword in s.title.lower():
            return s
    return None


class TestToXlsxStringOtherFossils:
    def test_single_string_row_rendered(self):
        data = _other_data("Genus sp. A")
        out = to_xlsx(data)
        wb = load_workbook(io.BytesIO(out))
        ws = _find_sheet(wb, "fossil") or wb.worksheets[-1]
        flat = " | ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
        assert "Genus sp. A" in flat, (
            f"string other_fossils entry missing from sheet {ws.title!r}: {flat!r}"
        )

    def test_multiple_string_rows_rendered(self):
        data = _other_data("Genus A sp.", "Genus B sp.", "Genus C sp.")
        out = to_xlsx(data)
        wb = load_workbook(io.BytesIO(out))
        ws = _find_sheet(wb, "fossil") or wb.worksheets[-1]
        flat = " | ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
        for name in ("Genus A sp.", "Genus B sp.", "Genus C sp."):
            assert name in flat, f"{name!r} not in xlsx output: {flat!r}"

    def test_empty_strings_skipped(self):
        data = _other_data("Genus A sp.", "", "  ", "Genus B sp.")
        out = to_xlsx(data)
        wb = load_workbook(io.BytesIO(out))
        ws = _find_sheet(wb, "fossil") or wb.worksheets[-1]
        flat = " | ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
        assert "Genus A sp." in flat
        assert "Genus B sp." in flat
        # Should not contain empty/whitespace standalone rows
        assert "  |" not in flat